from collections.abc import Generator
from datetime import date, datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.deps import get_db
from app.core.database import Base
from app.main import app
from app.models.intern_profile import InternProfile
from app.models.internship_program import InternshipProgram
from app.models.program_member import ProgramMember
from app.models.role import Role
from app.models.user import User
from app.utils.authenticate_login import create_access_token
import app.models as _models  # noqa: F401


@pytest.fixture
def analytics_client() -> Generator[tuple[TestClient, sessionmaker], None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_factory = sessionmaker(bind=engine, autoflush=False)
    Base.metadata.create_all(engine)

    with session_factory() as db:
        hr_role = Role(name="hr", description="HR")
        intern_role = Role(name="intern", description="Intern")
        admin_role = Role(name="admin", description="Admin")
        db.add_all([hr_role, intern_role, admin_role])
        db.flush()

        hr = User(
            id=1,
            code="HR0001",
            email="hr@example.com",
            password_hash="hashed",
            full_name="HR User",
            role_id=hr_role.id,
            status="active",
        )
        admin = User(
            id=2,
            code="AD0001",
            email="admin@example.com",
            password_hash="hashed",
            full_name="Admin User",
            role_id=admin_role.id,
            status="active",
        )
        intern1 = User(
            id=10,
            code="TTS0001",
            email="tts1@example.com",
            password_hash="hashed",
            full_name="TTS Mot",
            role_id=intern_role.id,
            status="active",
        )
        intern2 = User(
            id=11,
            code="TTS0002",
            email="tts2@example.com",
            password_hash="hashed",
            full_name="TTS Hai",
            role_id=intern_role.id,
            status="active",
        )
        intern3 = User(
            id=12,
            code="TTS0003",
            email="tts3@example.com",
            password_hash="hashed",
            full_name="TTS Ba",
            role_id=intern_role.id,
            status="active",
        )
        program = InternshipProgram(
            id=1,
            name="Dot 1",
            department="IT",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 6, 1),
            max_interns=50,
            status="open",
        )
        db.add_all([hr, admin, intern1, intern2, intern3, program])
        db.flush()

        # Intern profiles
        p1 = InternProfile(
            user_id=10,
            status="approved",
            university="ICTU",
            major="Cong nghe thong tin",
            created_at=datetime(2026, 3, 1, 10, 0, 0),
        )
        p2 = InternProfile(
            user_id=11,
            status="approved",
            university="ICTU",
            major="Ky thuat phan mem",
            created_at=datetime(2026, 3, 2, 10, 0, 0),
        )
        p3 = InternProfile(
            user_id=12,
            status="approved",
            university="Bach Khoa",
            major="Cong nghe thong tin",
            created_at=datetime(2026, 3, 3, 10, 0, 0),
        )
        db.add_all([p1, p2, p3])
        db.flush()

        # Program members: intern 10 & 11 thuộc program 1
        pm1 = ProgramMember(program_id=1, intern_user_id=10)
        pm2 = ProgramMember(program_id=1, intern_user_id=11)
        db.add_all([pm1, pm2])
        db.commit()

    def override_get_db() -> Generator[Session, None, None]:
        with session_factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    yield client, session_factory
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)
    engine.dispose()


def test_get_analytics_sources_success(analytics_client):
    client, _ = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    response = client.get(
        "/api/hr/analytics/sources",
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["total_interns"] == 3

    # Kiểm tra thống kê trường: ICTU có 2 (66.67%), Bach Khoa có 1 (33.33%)
    unis = {item["name"]: item for item in data["by_university"]}
    assert unis["ICTU"]["count"] == 2
    assert round(unis["ICTU"]["percentage"], 1) == 66.7
    assert unis["Bach Khoa"]["count"] == 1
    assert round(unis["Bach Khoa"]["percentage"], 1) == 33.3

    # Kiểm tra chuyên ngành: CNTT có 2, KTPM có 1
    majors = {item["name"]: item for item in data["by_major"]}
    assert majors["Cong nghe thong tin"]["count"] == 2
    assert majors["Ky thuat phan mem"]["count"] == 1


def test_get_analytics_sources_filter_program(analytics_client):
    client, _ = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    response = client.get(
        "/api/hr/analytics/sources?program_id=1",
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    data = response.json()
    # Program 1 chỉ có 2 interns (10 và 11), cả 2 đều từ ICTU
    assert data["total_interns"] == 2
    assert len(data["by_university"]) == 1
    assert data["by_university"][0]["name"] == "ICTU"
    assert data["by_university"][0]["count"] == 2
    assert data["by_university"][0]["percentage"] == 100.0


def test_get_analytics_sources_empty(analytics_client):
    client, _ = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    response = client.get(
        "/api/hr/analytics/sources?status=rejected",
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["total_interns"] == 0
    assert data["by_university"] == []
    assert data["by_major"] == []


def test_new_empty_program_analytics_returns_zero_without_division_by_zero(
    analytics_client,
):
    client, session_factory = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    with session_factory() as db:
        db.add(
            InternshipProgram(
                id=2,
                name="Dot moi chua co thuc tap sinh",
                department="IT",
                start_date=date(2026, 7, 1),
                end_date=date(2026, 12, 31),
                max_interns=20,
                status="open",
            )
        )
        db.commit()

    response = client.get(
        "/api/hr/analytics/sources?program_id=2",
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["total_interns"] == 0
    assert data["by_university"] == []
    assert data["by_major"] == []


def test_analytics_sources_with_one_intern_returns_100_percent(analytics_client):
    client, session_factory = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    with session_factory() as db:
        db.add(
            InternshipProgram(
                id=2,
                name="Dot moi co mot thuc tap sinh",
                department="IT",
                start_date=date(2026, 7, 1),
                end_date=date(2026, 12, 31),
                max_interns=20,
                status="open",
            )
        )
        db.add(ProgramMember(program_id=2, intern_user_id=10))
        db.commit()

    response = client.get(
        "/api/hr/analytics/sources?program_id=2",
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["total_interns"] == 1
    assert data["by_university"] == [
        {"name": "ICTU", "count": 1, "percentage": 100.0}
    ]
    assert data["by_major"] == [
        {"name": "Cong nghe thong tin", "count": 1, "percentage": 100.0}
    ]


def test_get_analytics_sources_forbidden_for_intern(analytics_client):
    client, _ = analytics_client
    intern_token = create_access_token(user_id=10, role="intern")

    response = client.get(
        "/api/hr/analytics/sources",
        headers={"Authorization": f"Bearer {intern_token}"},
    )
    assert response.status_code == 403


def test_export_analytics_xlsx_with_filters(analytics_client):
    client, _ = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    response = client.get(
        "/api/hr/analytics/export",
        params={
            "format": "xlsx",
            "program_id": 1,
            "status": "approved",
            "from_date": "2026-03-01",
            "to_date": "2026-03-02",
        },
        headers={"Authorization": f"Bearer {hr_token}"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="intern-analytics-')
    assert disposition.endswith('.xlsx"')

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["Tổng quan", "Theo trường", "Theo ngành"]
    summary = workbook["Tổng quan"]
    assert summary["A10"].value == "Tổng số thực tập sinh"
    assert summary["B10"].value == 2
    assert summary["B5"].value == "1"
    university_rows = list(workbook["Theo trường"].values)
    assert ("ICTU", 2, 1) in university_rows


def test_export_analytics_pdf_contains_vietnamese(analytics_client):
    client, session_factory = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")
    with session_factory() as db:
        profile = db.query(InternProfile).filter(InternProfile.user_id == 10).one()
        profile.university = "Đại học Thái Nguyên"
        db.commit()

    response = client.get(
        "/api/hr/analytics/export?format=pdf&program_id=1",
        headers={"Authorization": f"Bearer {hr_token}"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="intern-analytics-')
    assert disposition.endswith('.pdf"')
    assert response.content.startswith(b"%PDF-")
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages)
    assert "BÁO CÁO THỐNG KÊ THỰC TẬP SINH" in pdf_text
    assert "Đại học Thái Nguyên" in pdf_text
    assert "Tổng số thực tập sinh: 2" in pdf_text


@pytest.mark.parametrize("file_format", ["xlsx", "pdf"])
def test_export_analytics_empty_data(analytics_client, file_format):
    client, _ = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    response = client.get(
        "/api/hr/analytics/export",
        params={"format": file_format, "status": "rejected"},
        headers={"Authorization": f"Bearer {hr_token}"},
    )

    assert response.status_code == 200
    if file_format == "xlsx":
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        assert workbook["Tổng quan"]["B10"].value == 0
        assert workbook["Theo trường"]["A2"].value == "Không có dữ liệu"
        assert workbook["Theo ngành"]["A2"].value == "Không có dữ liệu"
    else:
        pdf_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(BytesIO(response.content)).pages
        )
        assert "Tổng số thực tập sinh: 0" in pdf_text
        assert pdf_text.count("Không có dữ liệu") == 2


def test_export_analytics_rejects_invalid_format(analytics_client):
    client, _ = analytics_client
    hr_token = create_access_token(user_id=1, role="hr")

    response = client.get(
        "/api/hr/analytics/export?format=csv",
        headers={"Authorization": f"Bearer {hr_token}"},
    )

    assert response.status_code == 422


def test_export_analytics_requires_hr_or_admin(analytics_client):
    client, _ = analytics_client
    intern_token = create_access_token(user_id=10, role="intern")

    unauthenticated = client.get("/api/hr/analytics/export?format=xlsx")
    forbidden = client.get(
        "/api/hr/analytics/export?format=xlsx",
        headers={"Authorization": f"Bearer {intern_token}"},
    )
    admin_token = create_access_token(user_id=2, role="admin")
    admin_response = client.get(
        "/api/hr/analytics/export?format=xlsx",
        headers={"Authorization": f"Bearer {admin_token}"},
    )

    assert unauthenticated.status_code == 401
    assert forbidden.status_code == 403
    assert admin_response.status_code == 200
