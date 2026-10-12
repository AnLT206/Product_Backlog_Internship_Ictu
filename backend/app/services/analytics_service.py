import os
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import and_, case, func
from sqlalchemy.orm import Session

from app.models.intern_profile import InternProfile
from app.models.internship_program import InternshipProgram
from app.models.program_member import ProgramMember
from app.schemas.analytics import (
    ComparisonKPI,
    CompletionRateAnalyticsResponse,
    InternSourceAnalyticsResponse,
    KPICardItem,
    PeriodSummaryItem,
    ProgramBatchCompletionItem,
    StatItem,
)


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"


@dataclass(frozen=True)
class AnalyticsExport:
    content: bytes
    filename: str
    media_type: str


class AnalyticsService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def get_sources_analytics(
        self,
        program_id: int | None = None,
        status: str | None = None,
        from_date: date | None = None,
        to_date: date | None = None,
    ) -> InternSourceAnalyticsResponse:
        """
        Viết câu truy vấn (COUNT, GROUP BY) thống kê số lượng và tỷ lệ phần trăm
        thực tập sinh theo từng Trường đại học và Chuyên ngành (SCRUM-176).
        """
        # Base query cho InternProfile
        base_query = self.db.query(InternProfile)

        if program_id is not None:
            base_query = base_query.join(
                ProgramMember, ProgramMember.intern_user_id == InternProfile.user_id
            ).filter(ProgramMember.program_id == program_id)

        if status:
            base_query = base_query.filter(InternProfile.status == status.strip())

        if from_date:
            from_dt = datetime.combine(from_date, datetime.min.time())
            base_query = base_query.filter(InternProfile.created_at >= from_dt)

        if to_date:
            to_dt = datetime.combine(to_date, datetime.max.time())
            base_query = base_query.filter(InternProfile.created_at <= to_dt)

        total_interns = base_query.count()

        if total_interns == 0:
            return InternSourceAnalyticsResponse(
                total_interns=0,
                by_university=[],
                by_major=[],
            )

        # 1. Thống kê theo Trường đại học (COUNT, GROUP BY university)
        uni_query = (
            base_query.with_entities(
                func.coalesce(InternProfile.university, "Khác").label("name"),
                func.count(InternProfile.id).label("count"),
            )
            .group_by(func.coalesce(InternProfile.university, "Khác"))
            .order_by(func.count(InternProfile.id).desc())
            .all()
        )

        by_university = [
            StatItem(
                name=row.name,
                count=row.count,
                percentage=round((row.count / total_interns) * 100, 2),
            )
            for row in uni_query
        ]

        # 2. Thống kê theo Chuyên ngành (COUNT, GROUP BY major)
        major_query = (
            base_query.with_entities(
                func.coalesce(InternProfile.major, "Khác").label("name"),
                func.count(InternProfile.id).label("count"),
            )
            .group_by(func.coalesce(InternProfile.major, "Khác"))
            .order_by(func.count(InternProfile.id).desc())
            .all()
        )

        by_major = [
            StatItem(
                name=row.name,
                count=row.count,
                percentage=round((row.count / total_interns) * 100, 2),
            )
            for row in major_query
        ]

        return InternSourceAnalyticsResponse(
            total_interns=total_interns,
            by_university=by_university,
            by_major=by_major,
        )

    def export_sources_report(
        self,
        *,
        file_format: str,
        program_id: int | None = None,
        status: str | None = None,
        from_date: date | None = None,
        to_date: date | None = None,
    ) -> AnalyticsExport:
        analytics = self.get_sources_analytics(
            program_id=program_id,
            status=status,
            from_date=from_date,
            to_date=to_date,
        )
        generated_at = datetime.now()
        filename = f"intern-analytics-{generated_at:%Y%m%d-%H%M%S}.{file_format}"
        filters = self._report_filters(program_id, status, from_date, to_date)

        if file_format == "xlsx":
            content = self._build_xlsx(analytics, filters, generated_at)
            media_type = XLSX_MEDIA_TYPE
        elif file_format == "pdf":
            content = self._build_pdf(analytics, filters, generated_at)
            media_type = PDF_MEDIA_TYPE
        else:
            raise ValueError("Định dạng báo cáo không được hỗ trợ.")

        return AnalyticsExport(content=content, filename=filename, media_type=media_type)

    @staticmethod
    def _report_filters(
        program_id: int | None,
        status: str | None,
        from_date: date | None,
        to_date: date | None,
    ) -> list[tuple[str, str]]:
        return [
            ("Chương trình", str(program_id) if program_id is not None else "Tất cả"),
            ("Trạng thái hồ sơ", status.strip() if status else "Tất cả"),
            ("Từ ngày", from_date.isoformat() if from_date else "Không giới hạn"),
            ("Đến ngày", to_date.isoformat() if to_date else "Không giới hạn"),
        ]

    def _build_xlsx(
        self,
        analytics: InternSourceAnalyticsResponse,
        filters: list[tuple[str, str]],
        generated_at: datetime,
    ) -> bytes:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Tổng quan"
        summary.merge_cells("A1:B1")
        summary["A1"] = "BÁO CÁO THỐNG KÊ THỰC TẬP SINH"
        summary["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        summary["A1"].fill = PatternFill("solid", fgColor="155EEF")
        summary["A1"].alignment = Alignment(horizontal="left")
        summary.append(["Ngày tạo", generated_at.strftime("%d/%m/%Y %H:%M")])
        summary.append([])
        summary.append(["Bộ lọc áp dụng", "Giá trị"])
        for row in summary.iter_rows(min_row=4, max_row=4):
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="344054")
        for filter_name, filter_value in filters:
            summary.append([filter_name, filter_value])
        summary.append([])
        summary.append(["Tổng số thực tập sinh", analytics.total_interns])
        summary.column_dimensions["A"].width = 34
        summary.column_dimensions["B"].width = 32

        self._add_xlsx_stats_sheet(workbook, "Theo trường", analytics.by_university)
        self._add_xlsx_stats_sheet(workbook, "Theo ngành", analytics.by_major)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _add_xlsx_stats_sheet(
        workbook: Workbook,
        title: str,
        items: list[StatItem],
    ) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append([title, "Số thực tập sinh", "Tỷ lệ"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="344054")
        if items:
            for item in items:
                sheet.append([item.name, item.count, item.percentage / 100])
        else:
            sheet.append(["Không có dữ liệu", 0, 0])
        for cell in sheet["C"][1:]:
            cell.number_format = "0.00%"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 22
        sheet.column_dimensions["C"].width = 16

    def _build_pdf(
        self,
        analytics: InternSourceAnalyticsResponse,
        filters: list[tuple[str, str]],
        generated_at: datetime,
    ) -> bytes:
        regular_font, bold_font = self._register_pdf_fonts()
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            fontName=bold_font,
            fontSize=17,
            leading=22,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#155EEF"),
        )
        normal_style = ParagraphStyle(
            "ReportNormal", parent=styles["BodyText"], fontName=regular_font, fontSize=9
        )
        heading_style = ParagraphStyle(
            "ReportHeading",
            parent=styles["Heading2"],
            fontName=bold_font,
            fontSize=12,
            textColor=colors.HexColor("#344054"),
        )

        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=16 * mm,
            bottomMargin=16 * mm,
            title="Báo cáo thống kê thực tập sinh",
        )
        story = [
            Paragraph("BÁO CÁO THỐNG KÊ THỰC TẬP SINH", title_style),
            Spacer(1, 4 * mm),
            Paragraph(
                f"Ngày tạo: {generated_at:%d/%m/%Y %H:%M}",
                normal_style,
            ),
            Spacer(1, 2 * mm),
            Paragraph("Bộ lọc áp dụng", heading_style),
            self._pdf_table(
                [[Paragraph("Bộ lọc", normal_style), Paragraph("Giá trị", normal_style)]]
                + [
                    [
                        Paragraph(escape(name), normal_style),
                        Paragraph(escape(value), normal_style),
                    ]
                    for name, value in filters
                ],
                [62 * mm, 190 * mm],
            ),
            Spacer(1, 3 * mm),
            Paragraph(
                f"Tổng số thực tập sinh: {analytics.total_interns}",
                heading_style,
            ),
            Spacer(1, 2 * mm),
            Paragraph("Thống kê theo trường", heading_style),
            self._pdf_stats_table(analytics.by_university, normal_style),
            Spacer(1, 4 * mm),
            Paragraph("Thống kê theo ngành", heading_style),
            self._pdf_stats_table(analytics.by_major, normal_style),
        ]
        document.build(story)
        return output.getvalue()

    @staticmethod
    def _pdf_table(data: list[list[object]], widths: list[float]) -> Table:
        table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAECF0")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D0D5DD")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        return table

    def _pdf_stats_table(
        self,
        items: list[StatItem],
        normal_style: ParagraphStyle,
    ) -> Table:
        data: list[list[object]] = [
            [
                Paragraph("Tên", normal_style),
                Paragraph("Số thực tập sinh", normal_style),
                Paragraph("Tỷ lệ", normal_style),
            ]
        ]
        if items:
            data.extend(
                [
                    Paragraph(escape(item.name), normal_style),
                    str(item.count),
                    f"{item.percentage:.2f}%",
                ]
                for item in items
            )
        else:
            data.append([Paragraph("Không có dữ liệu", normal_style), "0", "0.00%"])
        return self._pdf_table(data, [168 * mm, 46 * mm, 38 * mm])

    @staticmethod
    def _register_pdf_fonts() -> tuple[str, str]:
        regular_candidates = [
            os.getenv("REPORT_PDF_FONT_PATH"),
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            r"C:\Windows\Fonts\arial.ttf",
        ]
        bold_candidates = [
            os.getenv("REPORT_PDF_FONT_BOLD_PATH"),
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            r"C:\Windows\Fonts\arialbd.ttf",
        ]
        regular_path = next(
            (Path(candidate) for candidate in regular_candidates if candidate and Path(candidate).is_file()),
            None,
        )
        bold_path = next(
            (Path(candidate) for candidate in bold_candidates if candidate and Path(candidate).is_file()),
            None,
        )
        if regular_path is None or bold_path is None:
            raise RuntimeError(
                "PDF export requires a Unicode font. Install fonts-dejavu-core "
                "or configure REPORT_PDF_FONT_PATH and REPORT_PDF_FONT_BOLD_PATH."
            )

        regular_name = "AnalyticsReportSans"
        bold_name = "AnalyticsReportSans-Bold"
        registered = set(pdfmetrics.getRegisteredFontNames())
        if regular_name not in registered:
            pdfmetrics.registerFont(TTFont(regular_name, str(regular_path)))
        if bold_name not in registered:
            pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
        return regular_name, bold_name

    def get_completion_rate_analytics(
        self,
        program_id: int | None = None,
        department: str | None = None,
        from_date: date | None = None,
        to_date: date | None = None,
    ) -> CompletionRateAnalyticsResponse:
        """API trả về các chỉ số KPI và tỷ lệ % hoàn thành chương trình của từng đợt (US-33)."""
        target_program: InternshipProgram | None = None
        if program_id is not None:
            target_program = (
                self.db.query(InternshipProgram)
                .filter(InternshipProgram.id == program_id, InternshipProgram.is_deleted.is_(False))
                .first()
            )
            if not target_program:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Không tìm thấy kỳ thực tập có ID {program_id}.",
                )

        # 1. Base query cho summary
        summary_query = (
            self.db.query(
                func.count(ProgramMember.id).label("total"),
                func.count(case((ProgramMember.status == "in_progress", ProgramMember.id))).label("in_progress"),
                func.count(case((ProgramMember.status == "completed", ProgramMember.id))).label("completed"),
                func.count(case((ProgramMember.status == "dropped_out", ProgramMember.id))).label("dropped_out"),
                func.count(case((ProgramMember.status == "failed", ProgramMember.id))).label("failed"),
            )
            .join(
                InternshipProgram,
                and_(
                    InternshipProgram.id == ProgramMember.program_id,
                    InternshipProgram.is_deleted.is_(False),
                ),
            )
        )

        if program_id is not None:
            summary_query = summary_query.filter(ProgramMember.program_id == program_id)
        if department and department.strip():
            summary_query = summary_query.filter(InternshipProgram.department == department.strip())
        if from_date:
            summary_query = summary_query.filter(InternshipProgram.start_date >= from_date)
        if to_date:
            summary_query = summary_query.filter(InternshipProgram.start_date <= to_date)

        sum_row = summary_query.one()
        total = sum_row.total or 0
        in_progress = sum_row.in_progress or 0
        completed = sum_row.completed or 0
        dropped_out = sum_row.dropped_out or 0
        failed = sum_row.failed or 0
        not_completed = dropped_out + failed

        completion_rate = round((completed / total) * 100, 2) if total > 0 else 0.0
        dropout_rate = round((dropped_out / total) * 100, 2) if total > 0 else 0.0
        failure_rate = round((failed / total) * 100, 2) if total > 0 else 0.0
        in_progress_rate = round((in_progress / total) * 100, 2) if total > 0 else 0.0
        not_completed_rate = round((not_completed / total) * 100, 2) if total > 0 else 0.0

        summary = {
            "total": total,
            "in_progress": in_progress,
            "completed": completed,
            "dropped_out": dropped_out,
            "failed": failed,
            "not_completed": not_completed,
            "completion_rate": completion_rate,
            "dropout_rate": dropout_rate,
            "failure_rate": failure_rate,
            "in_progress_rate": in_progress_rate,
            "not_completed_rate": not_completed_rate,
        }

        # 2. Tạo danh sách thẻ KPI
        kpis = [
            KPICardItem(
                key="total",
                label="Tổng thực tập sinh",
                value=total,
                rate=100.0 if total > 0 else 0.0,
                hint="Toàn bộ TTS trong bộ lọc",
                tone="primary",
            ),
            KPICardItem(
                key="completed",
                label="Đã hoàn thành",
                value=completed,
                rate=completion_rate,
                hint="Kết thúc chương trình thành công",
                tone="success",
            ),
            KPICardItem(
                key="in_progress",
                label="Đang thực hiện",
                value=in_progress,
                rate=in_progress_rate,
                hint="Hiện đang trong chương trình",
                tone="warn",
            ),
            KPICardItem(
                key="not_completed",
                label="Không hoàn thành",
                value=not_completed,
                rate=not_completed_rate,
                hint="Bỏ cuộc hoặc không đạt",
                tone="danger",
            ),
        ]

        # 3. Danh sách tỷ lệ hoàn thành theo từng đợt (programs)
        prog_query = (
            self.db.query(
                InternshipProgram.id,
                InternshipProgram.name,
                InternshipProgram.department,
                InternshipProgram.start_date,
                InternshipProgram.end_date,
                InternshipProgram.status,
                func.count(ProgramMember.id).label("total"),
                func.count(case((ProgramMember.status == "in_progress", ProgramMember.id))).label("in_progress"),
                func.count(case((ProgramMember.status == "completed", ProgramMember.id))).label("completed"),
                func.count(case((ProgramMember.status == "dropped_out", ProgramMember.id))).label("dropped_out"),
                func.count(case((ProgramMember.status == "failed", ProgramMember.id))).label("failed"),
            )
            .outerjoin(
                ProgramMember,
                ProgramMember.program_id == InternshipProgram.id,
            )
            .filter(InternshipProgram.is_deleted.is_(False))
        )

        if program_id is not None:
            prog_query = prog_query.filter(InternshipProgram.id == program_id)
        if department and department.strip():
            prog_query = prog_query.filter(InternshipProgram.department == department.strip())
        if from_date:
            prog_query = prog_query.filter(InternshipProgram.start_date >= from_date)
        if to_date:
            prog_query = prog_query.filter(InternshipProgram.start_date <= to_date)

        prog_rows = (
            prog_query.group_by(
                InternshipProgram.id,
                InternshipProgram.name,
                InternshipProgram.department,
                InternshipProgram.start_date,
                InternshipProgram.end_date,
                InternshipProgram.status,
            )
            .order_by(
                InternshipProgram.start_date.desc(),
                InternshipProgram.id.desc(),
            )
            .all()
        )

        programs_list: list[ProgramBatchCompletionItem] = []
        for p in prog_rows:
            p_total = p.total or 0
            p_comp = p.completed or 0
            p_inp = p.in_progress or 0
            p_drop = p.dropped_out or 0
            p_fail = p.failed or 0
            p_not_comp = p_drop + p_fail
            p_rate = round((p_comp / p_total) * 100, 2) if p_total > 0 else 0.0

            programs_list.append(
                ProgramBatchCompletionItem(
                    id=p.id,
                    name=p.name,
                    department=p.department,
                    start_date=str(p.start_date) if p.start_date else None,
                    end_date=str(p.end_date) if p.end_date else None,
                    status=p.status or "open",
                    total=p_total,
                    completed=p_comp,
                    in_progress=p_inp,
                    dropped_out=p_drop,
                    failed=p_fail,
                    not_completed=p_not_comp,
                    completion_rate=p_rate,
                )
            )

        # 4. Tính toán kỳ trước và đối tượng so sánh
        previous_period: PeriodSummaryItem | None = None
        if target_program is not None:
            prev_prog = (
                self.db.query(InternshipProgram)
                .filter(
                    InternshipProgram.is_deleted.is_(False),
                    InternshipProgram.id != target_program.id,
                    InternshipProgram.start_date <= target_program.start_date,
                )
                .order_by(InternshipProgram.start_date.desc(), InternshipProgram.id.desc())
                .first()
            )
        else:
            prev_prog = None
            if len(prog_rows) >= 2:
                prev_prog_id = prog_rows[1].id
                prev_prog = (
                    self.db.query(InternshipProgram)
                    .filter(InternshipProgram.id == prev_prog_id)
                    .first()
                )

        if prev_prog:
            prev_row = (
                self.db.query(
                    func.count(ProgramMember.id).label("total"),
                    func.count(case((ProgramMember.status == "completed", ProgramMember.id))).label("completed"),
                )
                .filter(ProgramMember.program_id == prev_prog.id)
                .one()
            )
            prev_total = prev_row.total or 0
            prev_comp = prev_row.completed or 0
            prev_rate = round((prev_comp / prev_total) * 100, 2) if prev_total > 0 else 0.0

            previous_period = PeriodSummaryItem(
                id=prev_prog.id,
                name=prev_prog.name,
                total=prev_total,
                completed=prev_comp,
                completion_rate=prev_rate,
            )

        # So sánh chênh lệch
        if previous_period is not None:
            prev_r = previous_period.completion_rate
            diff = round(completion_rate - prev_r, 2)
            if diff > 0:
                trend = "up"
                arrow = "↑"
                badge_text = f"+{diff}%"
                text = f"↑ +{diff}% so với đợt trước"
            elif diff < 0:
                trend = "down"
                arrow = "↓"
                badge_text = f"{diff}%"
                text = f"↓ {diff}% so với đợt trước"
            else:
                trend = "same"
                arrow = "→"
                badge_text = "0%"
                text = "→ 0% so với đợt trước"

            comparison = ComparisonKPI(
                has_previous=True,
                current_rate=completion_rate,
                previous_rate=prev_r,
                diff=diff,
                trend=trend,
                arrow=arrow,
                text=text,
                badge_text=badge_text,
            )
        else:
            comparison = ComparisonKPI(
                has_previous=False,
                current_rate=completion_rate,
                previous_rate=0.0,
                diff=0.0,
                trend="same",
                arrow="→",
                text="Chưa có dữ liệu so sánh",
                badge_text="Chưa có dữ liệu so sánh",
            )

        return CompletionRateAnalyticsResponse(
            summary=summary,
            kpis=kpis,
            comparison=comparison,
            previous_period=previous_period,
            programs=programs_list,
        )

